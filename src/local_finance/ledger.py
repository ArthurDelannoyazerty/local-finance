from __future__ import annotations

import csv
import io
import sqlite3
import uuid
from collections import defaultdict
from collections.abc import Iterable
from datetime import UTC, date, datetime
from typing import Any, Literal

from openpyxl import Workbook

from .db import Database, database, row_to_dict, utc_now
from .schemas import AccountCreate, AccountUpdate, TradeInput, TradeUpdate


class RevisionConflict(RuntimeError):
    pass


class InventoryError(ValueError):
    pass


def _where_clause(
    *,
    start: date | None = None,
    end: date | None = None,
    transaction_type: str | None = None,
    account: str | None = None,
    category: str | None = None,
    query: str | None = None,
    visible_only: bool = False,
) -> tuple[str, list[Any]]:
    conditions = ["is_excluded = 0"]
    params: list[Any] = []
    if start:
        conditions.append("date >= ?")
        params.append(start.isoformat())
    if end:
        conditions.append("date <= ?")
        params.append(end.isoformat())
    if transaction_type:
        conditions.append("type = ?")
        params.append(transaction_type)
    if account:
        conditions.append("account = ?")
        params.append(account)
    if category:
        conditions.append("category = ?")
        params.append(category)
    if query:
        conditions.append(
            "(LOWER(category) LIKE ? OR LOWER(account) LIKE ? OR LOWER(comment) LIKE ? "
            "OR date LIKE ? OR CAST(amount AS TEXT) LIKE ?)"
        )
        term = f"%{query.strip().lower()}%"
        params.extend([term, term, term, term, term])
    if visible_only:
        conditions.append("account IN (SELECT name FROM accounts WHERE is_visible = 1)")
    return " AND ".join(conditions), params


def list_transactions(
    *,
    start: date | None = None,
    end: date | None = None,
    transaction_type: str | None = None,
    account: str | None = None,
    category: str | None = None,
    query: str | None = None,
    page: int = 1,
    page_size: int = 50,
    sort: str = "date",
    direction: Literal["asc", "desc"] = "desc",
    db: Database = database,
) -> dict[str, Any]:
    where, params = _where_clause(
        start=start,
        end=end,
        transaction_type=transaction_type,
        account=account,
        category=category,
        query=query,
    )
    sort_column = sort if sort in {"date", "amount", "category", "account", "type"} else "date"
    page = max(1, page)
    page_size = min(200, max(1, page_size))
    offset = (page - 1) * page_size
    with db.read() as connection:
        total = connection.execute(
            f"SELECT COUNT(*) FROM transactions WHERE {where}", params
        ).fetchone()[0]
        rows = connection.execute(
            f"""
            SELECT id, date, category, account, amount, currency, comment, type,
                   imported_at
            FROM transactions
            WHERE {where}
            ORDER BY {sort_column} {direction.upper()}, id {direction.upper()}
            LIMIT ? OFFSET ?
            """,
            [*params, page_size, offset],
        ).fetchall()
        totals = connection.execute(
            f"""
            SELECT
                COALESCE(SUM(CASE WHEN type = 'INCOME' THEN amount ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN type = 'EXPENSE' THEN amount ELSE 0 END), 0)
            FROM transactions WHERE {where}
            """,
            params,
        ).fetchone()
        categories = [
            row[0]
            for row in connection.execute(
                "SELECT DISTINCT category FROM transactions ORDER BY category"
            )
        ]
        accounts = [
            row[0]
            for row in connection.execute(
                "SELECT DISTINCT account FROM transactions ORDER BY account"
            )
        ]
    income, expense = float(totals[0]), float(totals[1])
    return {
        "items": [row_to_dict(row) for row in rows],
        "total": int(total),
        "page": page,
        "page_size": page_size,
        "summary": {"income": income, "expense": expense, "net": income - expense},
        "filters": {"categories": categories, "accounts": accounts},
    }


def get_date_bounds(*, db: Database = database) -> dict[str, str]:
    with db.read() as connection:
        row = connection.execute(
            """
            SELECT MIN(value), MAX(value) FROM (
                SELECT date AS value FROM transactions
                UNION ALL SELECT date FROM transfers
                UNION ALL SELECT date FROM investments
            )
            """
        ).fetchone()
    today = datetime.now(UTC).date().isoformat()
    return {"min": row[0] or today, "max": row[1] or today, "today": today}


def dashboard_summary(start: date, end: date, *, db: Database = database) -> dict[str, Any]:
    where, params = _where_clause(start=start, end=end, visible_only=True)
    with db.read() as connection:
        totals = connection.execute(
            f"""
            SELECT
                COALESCE(SUM(CASE WHEN type = 'INCOME' THEN amount ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN type = 'EXPENSE' THEN amount ELSE 0 END), 0)
            FROM transactions WHERE {where}
            """,
            params,
        ).fetchone()
        expenses = connection.execute(
            f"""
            SELECT category, SUM(amount) AS amount
            FROM transactions
            WHERE {where} AND type = 'EXPENSE'
            GROUP BY category ORDER BY amount DESC
            """,
            params,
        ).fetchall()
        monthly = connection.execute(
            f"""
            SELECT substr(date, 1, 7) AS month, type, SUM(amount) AS amount
            FROM transactions WHERE {where}
            GROUP BY month, type ORDER BY month
            """,
            params,
        ).fetchall()

    income, expense = float(totals[0]), float(totals[1])
    net = income - expense
    month_count = max(1, (end.year - start.year) * 12 + end.month - start.month + 1)
    monthly_map: dict[str, dict[str, float]] = defaultdict(lambda: {"income": 0.0, "expense": 0.0})
    for row in monthly:
        monthly_map[row["month"]][str(row["type"]).lower()] = float(row["amount"])
    trend: list[dict[str, Any]] = []
    expense_history: list[float] = []
    for month, values in monthly_map.items():
        expense_history.append(values["expense"])
        rolling = sum(expense_history[-3:]) / len(expense_history[-3:])
        trend.append({"month": month, **values, "expense_average_3m": rolling})
    return {
        "range": {"start": start.isoformat(), "end": end.isoformat()},
        "totals": {
            "income": income,
            "expense": expense,
            "savings": net,
            "savings_rate": (net / income * 100) if income else 0,
        },
        "monthly_average": {
            "income": income / month_count,
            "expense": expense / month_count,
            "savings": net / month_count,
        },
        "expenses_by_category": [row_to_dict(row) for row in expenses],
        "monthly": trend,
    }


def _sankey_payload(
    links: Iterable[tuple[str, str, float]],
) -> dict[str, list[dict[str, Any]]]:
    aggregated: dict[tuple[str, str], float] = defaultdict(float)
    for source, target, value in links:
        if source != target and value > 0:
            aggregated[(source, target)] += float(value)
    names = sorted({name for pair in aggregated for name in pair})
    return {
        "nodes": [{"name": name} for name in names],
        "links": [
            {"source": source, "target": target, "value": value}
            for (source, target), value in aggregated.items()
        ],
    }


def flow_summary(months: list[str], *, db: Database = database) -> dict[str, Any]:
    valid_months = sorted({month for month in months if len(month) == 7})
    if not valid_months:
        return {
            "cash_flow": {"nodes": [], "links": []},
            "transfers": {"nodes": [], "links": []},
            "investments": {"nodes": [], "links": []},
        }
    placeholders = ",".join("?" for _ in valid_months)
    with db.read() as connection:
        transactions = connection.execute(
            f"""
            SELECT type, category, SUM(amount) AS amount
            FROM transactions
            WHERE substr(date, 1, 7) IN ({placeholders}) AND is_excluded = 0
              AND account IN (SELECT name FROM accounts WHERE is_visible = 1)
            GROUP BY type, category
            """,
            valid_months,
        ).fetchall()
        transfers = connection.execute(
            f"""
            SELECT source_account, target_account, SUM(amount) AS amount
            FROM transfers WHERE substr(date, 1, 7) IN ({placeholders})
            GROUP BY source_account, target_account
            """,
            valid_months,
        ).fetchall()
        investments = connection.execute(
            f"""
            SELECT action, account, ticker,
                   SUM(quantity * unit_price + CASE WHEN action = 'BUY' THEN fees ELSE -fees END)
                       AS amount
            FROM investments WHERE substr(date, 1, 7) IN ({placeholders})
            GROUP BY action, account, ticker
            """,
            valid_months,
        ).fetchall()

    incomes = [
        (row["category"], float(row["amount"])) for row in transactions if row["type"] == "INCOME"
    ]
    expenses = [
        (row["category"], float(row["amount"])) for row in transactions if row["type"] == "EXPENSE"
    ]
    total_income = sum(value for _, value in incomes)
    total_expense = sum(value for _, value in expenses)
    cash_links: list[tuple[str, str, float]] = [
        (f"Revenu · {category}", "Total revenus", value) for category, value in incomes
    ]
    cash_links.extend(
        ("Total dépenses", f"Dépense · {category}", value) for category, value in expenses
    )
    common = min(total_income, total_expense)
    if common:
        cash_links.append(("Total revenus", "Total dépenses", common))
    if total_income > total_expense:
        cash_links.append(("Total revenus", "Épargne", total_income - total_expense))
    elif total_expense > total_income:
        cash_links.append(("Déficit", "Total dépenses", total_expense - total_income))

    transfer_links = [
        (str(row["source_account"]), str(row["target_account"]), float(row["amount"]))
        for row in transfers
    ]
    investment_links = [
        (
            str(row["account"]) if row["action"] == "BUY" else str(row["ticker"]),
            str(row["ticker"]) if row["action"] == "BUY" else str(row["account"]),
            max(0.0, float(row["amount"])),
        )
        for row in investments
    ]
    return {
        "cash_flow": _sankey_payload(cash_links),
        "transfers": _sankey_payload(transfer_links),
        "investments": _sankey_payload(investment_links),
    }


def list_accounts(*, db: Database = database) -> list[dict[str, Any]]:
    with db.read() as connection:
        rows = connection.execute(
            """
            SELECT name, initial_balance, opening_balance_date, is_visible,
                   revision, updated_at
            FROM accounts ORDER BY is_visible DESC, name
            """
        ).fetchall()
    return [{**row_to_dict(row), "is_visible": bool(row["is_visible"])} for row in rows]


def _validate_opening_date(
    connection: sqlite3.Connection,
    account: str,
    opening_date: date | None,
) -> None:
    if opening_date is None:
        return
    first_event = connection.execute(
        """
        SELECT MIN(event_date) FROM (
            SELECT date AS event_date FROM transactions WHERE account = ?
            UNION ALL SELECT date FROM investments WHERE account = ?
            UNION ALL SELECT date FROM transfers
                WHERE source_account = ? OR target_account = ?
        )
        """,
        (account, account, account, account),
    ).fetchone()[0]
    if first_event and opening_date.isoformat() > str(first_event):
        raise ValueError(
            f"Opening balance date cannot be later than the first {account} operation ({first_event})"
        )


def create_account(payload: AccountCreate, *, db: Database = database) -> dict[str, Any]:
    now = utc_now()
    with db.transaction(immediate=True) as connection:
        _validate_opening_date(connection, payload.name, payload.opening_balance_date)
        connection.execute(
            """
            INSERT INTO accounts(
                name, initial_balance, opening_balance_date, is_visible,
                revision, updated_at
            ) VALUES (?, ?, ?, ?, 1, ?)
            """,
            (
                payload.name,
                payload.initial_balance,
                payload.opening_balance_date.isoformat() if payload.opening_balance_date else None,
                int(payload.is_visible),
                now,
            ),
        )
    return {**payload.model_dump(mode="json"), "revision": 1, "updated_at": now}


def update_account(
    name: str,
    payload: AccountUpdate,
    *,
    db: Database = database,
) -> dict[str, Any]:
    now = utc_now()
    with db.transaction(immediate=True) as connection:
        current = connection.execute(
            "SELECT revision FROM accounts WHERE name = ?", (name,)
        ).fetchone()
        if current is None:
            raise KeyError("Account not found")
        if current["revision"] != payload.revision:
            raise RevisionConflict("The account changed in another browser tab")
        _validate_opening_date(connection, name, payload.opening_balance_date)
        cursor = connection.execute(
            """
            UPDATE accounts
            SET initial_balance = ?, opening_balance_date = ?, is_visible = ?,
                revision = revision + 1, updated_at = ?
            WHERE name = ? AND revision = ?
            """,
            (
                payload.initial_balance,
                payload.opening_balance_date.isoformat() if payload.opening_balance_date else None,
                int(payload.is_visible),
                now,
                name,
                payload.revision,
            ),
        )
        if cursor.rowcount != 1:
            raise RevisionConflict("The account changed in another browser tab")
    return {
        "name": name,
        **payload.model_dump(mode="json", exclude={"revision"}),
        "revision": payload.revision + 1,
        "updated_at": now,
    }


def _trade_filters(
    *,
    start: date | None,
    end: date | None,
    action: str | None,
    account: str | None,
    ticker: str | None,
    query: str | None,
) -> tuple[str, list[Any]]:
    conditions = ["1 = 1"]
    params: list[Any] = []
    for expression, value in (
        ("date >= ?", start.isoformat() if start else None),
        ("date <= ?", end.isoformat() if end else None),
        ("action = ?", action),
        ("account = ?", account),
        ("ticker = ?", ticker.upper() if ticker else None),
    ):
        if value is not None:
            conditions.append(expression)
            params.append(value)
    if query:
        term = f"%{query.strip().lower()}%"
        conditions.append(
            "(LOWER(ticker) LIKE ? OR LOWER(name) LIKE ? OR LOWER(account) LIKE ? "
            "OR LOWER(comment) LIKE ? OR date LIKE ? OR CAST(quantity AS TEXT) LIKE ? "
            "OR CAST(unit_price AS TEXT) LIKE ?)"
        )
        params.extend([term, term, term, term, term, term, term])
    return " AND ".join(conditions), params


def list_trades(
    *,
    start: date | None = None,
    end: date | None = None,
    action: str | None = None,
    account: str | None = None,
    ticker: str | None = None,
    query: str | None = None,
    page: int = 1,
    page_size: int = 50,
    db: Database = database,
) -> dict[str, Any]:
    where, params = _trade_filters(
        start=start, end=end, action=action, account=account, ticker=ticker, query=query
    )
    page, page_size = max(1, page), min(200, max(1, page_size))
    with db.read() as connection:
        total = connection.execute(
            f"SELECT COUNT(*) FROM investments WHERE {where}", params
        ).fetchone()[0]
        rows = connection.execute(
            f"""
            SELECT id, date, ticker, name, action, quantity, unit_price, fees,
                   currency, account, comment, revision, created_at, updated_at
            FROM investments WHERE {where}
            ORDER BY date DESC, created_at DESC, id DESC
            LIMIT ? OFFSET ?
            """,
            [*params, page_size, (page - 1) * page_size],
        ).fetchall()
        selectors = {
            "accounts": [
                row[0]
                for row in connection.execute(
                    "SELECT name FROM accounts WHERE is_visible = 1 ORDER BY name"
                )
            ],
            "tickers": [
                row[0]
                for row in connection.execute(
                    "SELECT DISTINCT ticker FROM investments ORDER BY ticker"
                )
            ],
        }
    return {
        "items": [row_to_dict(row) for row in rows],
        "total": int(total),
        "page": page,
        "page_size": page_size,
        "filters": selectors,
    }


def _assert_inventory(
    connection: sqlite3.Connection,
    payload: TradeInput,
    *,
    exclude_id: str | None = None,
) -> None:
    rows = connection.execute(
        """
        SELECT id, date, action, quantity FROM investments
        WHERE account = ? AND ticker = ? AND (? IS NULL OR id != ?)
        """,
        (payload.account, payload.ticker, exclude_id, exclude_id),
    ).fetchall()
    ledger = [
        (
            str(row["date"]),
            0 if row["action"] == "BUY" else 1,
            str(row["id"]),
            str(row["action"]),
            float(row["quantity"]),
        )
        for row in rows
    ]
    ledger.append(
        (
            payload.date.isoformat(),
            0 if payload.action == "BUY" else 1,
            exclude_id or "~new",
            payload.action,
            payload.quantity,
        )
    )
    quantity = 0.0
    for trade_date, _, _, action, trade_quantity in sorted(ledger):
        quantity += trade_quantity if action == "BUY" else -trade_quantity
        if quantity < -1e-9:
            raise InventoryError(
                f"This trade would make {payload.ticker} holdings negative on {trade_date}"
            )


def create_trade(payload: TradeInput, *, db: Database = database) -> dict[str, Any]:
    trade_id, now = str(uuid.uuid4()), utc_now()
    with db.transaction(immediate=True) as connection:
        account_exists = connection.execute(
            "SELECT 1 FROM accounts WHERE name = ?", (payload.account,)
        ).fetchone()
        if not account_exists:
            raise ValueError("The selected account does not exist")
        _assert_inventory(connection, payload)
        connection.execute(
            """
            INSERT INTO investments(
                id, date, ticker, name, action, quantity, unit_price, fees,
                currency, account, comment, revision, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
            """,
            (
                trade_id,
                payload.date.isoformat(),
                payload.ticker,
                payload.name,
                payload.action,
                payload.quantity,
                payload.unit_price,
                payload.fees,
                payload.currency,
                payload.account,
                payload.comment,
                now,
                now,
            ),
        )
    return {
        "id": trade_id,
        **payload.model_dump(mode="json"),
        "revision": 1,
        "created_at": now,
        "updated_at": now,
    }


def update_trade(
    trade_id: str,
    payload: TradeUpdate,
    *,
    db: Database = database,
) -> dict[str, Any]:
    now = utc_now()
    with db.transaction(immediate=True) as connection:
        current = connection.execute(
            "SELECT revision FROM investments WHERE id = ?", (trade_id,)
        ).fetchone()
        if current is None:
            raise KeyError("Trade not found")
        if current["revision"] != payload.revision:
            raise RevisionConflict("The trade changed in another browser tab")
        account_exists = connection.execute(
            "SELECT 1 FROM accounts WHERE name = ?", (payload.account,)
        ).fetchone()
        if not account_exists:
            raise ValueError("The selected account does not exist")
        _assert_inventory(connection, payload, exclude_id=trade_id)
        cursor = connection.execute(
            """
            UPDATE investments
            SET date = ?, ticker = ?, name = ?, action = ?, quantity = ?,
                unit_price = ?, fees = ?, currency = ?, account = ?, comment = ?,
                revision = revision + 1, updated_at = ?
            WHERE id = ? AND revision = ?
            """,
            (
                payload.date.isoformat(),
                payload.ticker,
                payload.name,
                payload.action,
                payload.quantity,
                payload.unit_price,
                payload.fees,
                payload.currency,
                payload.account,
                payload.comment,
                now,
                trade_id,
                payload.revision,
            ),
        )
        if cursor.rowcount != 1:
            raise RevisionConflict("The trade changed in another browser tab")
    return {
        "id": trade_id,
        **payload.model_dump(mode="json", exclude={"revision"}),
        "revision": payload.revision + 1,
        "updated_at": now,
    }


def delete_trade(
    trade_id: str,
    revision: int,
    *,
    db: Database = database,
) -> None:
    with db.transaction(immediate=True) as connection:
        row = connection.execute(
            "SELECT * FROM investments WHERE id = ? AND revision = ?",
            (trade_id, revision),
        ).fetchone()
        if row is None:
            raise RevisionConflict("The trade changed or was already deleted")
        connection.execute("DELETE FROM investments WHERE id = ?", (trade_id,))
        remaining = connection.execute(
            """
            SELECT date, action, quantity FROM investments
            WHERE account = ? AND ticker = ? ORDER BY date, action
            """,
            (row["account"], row["ticker"]),
        ).fetchall()
        quantity = 0.0
        for item in remaining:
            quantity += (
                float(item["quantity"]) if item["action"] == "BUY" else -float(item["quantity"])
            )
            if quantity < -1e-9:
                raise InventoryError("Deleting this buy would leave a later sale without inventory")


def _spreadsheet_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Export"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column in worksheet.columns:
        width = min(42, max(12, *(len(str(cell.value or "")) + 2 for cell in column)))
        worksheet.column_dimensions[column[0].column_letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_trades(
    *,
    file_format: Literal["csv", "xlsx"],
    start: date | None = None,
    end: date | None = None,
    action: str | None = None,
    account: str | None = None,
    ticker: str | None = None,
    query: str | None = None,
    db: Database = database,
) -> bytes:
    where, params = _trade_filters(
        start=start, end=end, action=action, account=account, ticker=ticker, query=query
    )
    headers = [
        "Date",
        "Action",
        "Ticker",
        "Nom",
        "Quantité",
        "Prix unitaire",
        "Frais",
        "Devise",
        "Compte",
        "Commentaire",
    ]
    with db.read() as connection:
        records = connection.execute(
            f"""
            SELECT date, action, ticker, name, quantity, unit_price, fees,
                   currency, account, comment
            FROM investments WHERE {where} ORDER BY date DESC, id DESC
            """,
            params,
        ).fetchall()
    rows = [list(record) for record in records]
    if file_format == "xlsx":
        return _spreadsheet_bytes(headers, rows)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")
